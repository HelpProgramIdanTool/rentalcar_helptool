"""Read-only spreadsheet extraction. Column mappings are supplied by the operator."""
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX as COLOR_INDEXED
from openpyxl.utils.cell import column_index_from_string


def identifier(value):
    value = str(value or '').replace('\u200b','').replace('\ufeff','').strip().upper()
    if re.fullmatch(r'\d+\.0', value):
        value = value[:-2]
    return re.sub(r'^(?:[\[(]?RC[\])]?|NN|HT)\s*(?=\d+$)', '', value)


def amount(value):
    if value in (None, ''):
        return None
    try:
        result = Decimal(str(value).replace('\xa0', '').replace(' ', '').replace(',', '.'))
        if not result.is_finite():
            raise ValueError('Некорректная сумма')
        return str(result.quantize(Decimal('.01')))
    except InvalidOperation:
        raise ValueError(f'Не удалось прочитать сумму: {str(value)[:40]}')


def date_value(value):
    if isinstance(value, (date, datetime)):
        return value.strftime('%Y-%m-%d')
    value = str(value or '')
    try:
        return date.fromisoformat(value[:10]).isoformat()
    except ValueError:
        pass
    match = re.search(r'(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})', value)
    if match:
        day, month, year = map(int, match.groups())
        try:
            return date(year + 2000 if year < 100 else year, month, day).isoformat()
        except ValueError:
            return ''
    try:
        return date.fromisoformat(value[:10]).isoformat()
    except ValueError:
        return ''


def color(value):
    if value is None:
        return ''
    if value.type == 'rgb':
        return value.rgb[-6:].upper()
    if value.type == 'indexed' and value.indexed < len(COLOR_INDEXED):
        return COLOR_INDEXED[value.indexed][-6:].upper()
    return ''


def read_rows(content, mapping, *, sheet='', first_row=2, color_rules=False):
    if len(content) > 20 * 1024 * 1024:
        raise ValueError('Максимальный размер файла — 20 МБ.')
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 120 * 1024 * 1024:
                raise ValueError('Слишком большой распакованный файл.')
    except BadZipFile:
        raise ValueError('Нужен файл Excel .xlsx.')
    book = load_workbook(BytesIO(content), data_only=True)
    ws = book[sheet] if sheet else book.active
    if ws.max_row > 30000 or ws.max_column > 300:
        raise ValueError('Слишком большая таблица. Выберите файл только с заказами.')
    indexes = {key: column_index_from_string(col.upper()) for key, col in mapping.items() if col}
    if 'number' not in indexes:
        raise ValueError('Укажите столбец номера заказа.')
    results = []
    for cells in ws.iter_rows(min_row=first_row):
        def get(key):
            idx = indexes.get(key, 0)
            return cells[idx-1].value if idx and idx <= len(cells) else None
        num = identifier(get('number'))
        name = str(get('name') or '').strip()
        # Header/summary rows must not become reservations.
        if not num and not (name and (get('pickup') or get('return'))):
            continue
        if num and not re.search(r'\d', num):
            if name and (get('pickup') or get('return')):
                num = ''
            else:
                continue
        row = {'number': num, 'name': name, 'row': cells[0].row}
        row['alternate_number'] = identifier(get('alternate_number'))
        for key in ('gross', 'net', 'amount'):
            row[key] = amount(get(key))
        for key in ('pickup', 'return'):
            row[key] = date_value(get(key))
            row[key + '_raw'] = str(get(key) or '')
        for key in ('email', 'phone', 'invoice', 'orderer'):
            row[key] = str(get(key) or '').strip()
        red = any(color(c.font.color) in ('FF0000', 'EE0000') or (c.fill.patternType and color(c.fill.fgColor) in ('FF0000', 'EE0000')) for c in cells[:9])
        row['cancelled'] = bool(color_rules and (red or row['amount'] is None))
        row['paid_color'] = bool(color_rules and any(c.fill.patternType and color(c.fill.fgColor) not in ('', 'FFFFFF', '000000', 'FF0000', 'EE0000') for c in cells[:9]))
        results.append(row)
    book.close()
    if not results:
        raise ValueError('Не найдено заказов. Проверьте лист и столбцы.')
    return results
