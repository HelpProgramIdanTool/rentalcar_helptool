import hashlib
import tempfile
from uuid import uuid4
from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from pathlib import Path
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from bookings.models import Booking
from customers.models import Customer, CustomerEvent
from suppliers.models import Supplier
from .models import Reconciliation, Row, Revision
from .parsing import read_rows, identifier, date_value
from .services import row_net, money, totals, stage, import_own, decide, close, checkpoint, workbook


class CalculationTests(SimpleTestCase):
    def test_prefixes_and_iso_dates(self):
        for text in ['RC123', '(RC) 123', '[RC]123', 'NN 123', 'HT 123']:
            self.assertEqual(identifier(text), '123')
        self.assertEqual(identifier('26-0002-7008'), '26-0002-7008')
        self.assertEqual(date_value('2026-09-16'), '2026-09-16')

    def test_supplier_net_wins_over_gross(self):
        row = SimpleNamespace(supplier_lines=[{'net':'100.00','gross':'900.00'}, {'net':'50.00'}])
        self.assertEqual(row_net(row, 23), Decimal('150.00'))

    def test_gross_needs_confirmed_vat(self):
        row = SimpleNamespace(supplier_lines=[{'gross':'123.00'}])
        self.assertIsNone(row_net(row, None))
        self.assertEqual(row_net(row, Decimal(23)), Decimal('100.00'))

    def test_commission_half_up(self):
        self.assertEqual(money(Decimal('25.125') * Decimal('.04')), Decimal('1.01'))

    def test_excel_never_executes_source_text(self):
        data = workbook({'summary': {'Invoice':'=HYPERLINK("bad")'}, 'rows':[]})
        book = load_workbook(BytesIO(data), data_only=False)
        self.assertEqual(book.active['B1'].data_type, 's')

    def test_parser_mapping(self):
        book = Workbook();ws = book.active
        ws.append(['Number','Net','Name']);ws.append(['0099', 100, 'TEST-NAME'])
        data = BytesIO();book.save(data)
        rows = read_rows(data.getvalue(), {'number':'A','net':'B','name':'C'})
        self.assertEqual(rows[0]['number'], '0099')
        self.assertEqual(rows[0]['net'], '100.00')
        self.assertEqual(rows[0]['name'], 'TEST-NAME')


class WorkflowTests(TestCase):
    def setUp(self):
        root = Path('.test-temp').resolve()
        root.mkdir(exist_ok=True)
        directory = root / ('reconciliation-' + uuid4().hex)
        directory.mkdir()
        self.override = override_settings(RECONCILIATION_EXPORT_ROOT=directory)
        self.override.enable();self.addCleanup(self.override.disable)
        self.supplier = Supplier.objects.create(supplier_code='TEST', supplier_name='TEST-SUPPLIER')
        self.customer = Customer.objects.create(first_name='TEST',last_name='CLIENT',full_name_latin='TEST CLIENT',phone_1='1234567')
        self.booking = Booking.objects.create(customer=self.customer,supplier=self.supplier,supplier_booking_number='123',status='COMPLETED',source_quote_snapshot={'historical_import':{'test':True}},total_price_gross=123,vehicle_price_gross=123)
        self.batch = Reconciliation.objects.create(month=date(2026,9,1),supplier=self.supplier,vat=23)
        self.actor = get_user_model().objects.create_superuser('reviewer','reviewer@example.test','test-password')

    def row(self, **kwargs):
        defaults = dict(reconciliation=self.batch,number='123',booking=self.booking,own={'name':'TEST CLIENT','amount':'123.00'},supplier_lines=[{'net':'100.00'}],employee_data={'amount':'9000'},decision='INCLUDE')
        defaults.update(kwargs)
        return Row.objects.create(**defaults)

    def test_duplicate_upload_and_replacement(self):
        rows=[{'number':'123','name':'TEST CLIENT','net':'100.00','row':2}]
        self.assertTrue(stage(self.batch,'SUPPLIER','x.xlsx',b'first',{},rows,self.actor))
        self.assertFalse(stage(self.batch,'SUPPLIER','renamed.xlsx',b'first',{},rows,self.actor))
        stage(self.batch,'SUPPLIER','new.xlsx',b'second',{},[dict(rows[0],net='200.00')],self.actor)
        self.assertEqual(self.batch.rows.count(),1)
        self.assertEqual(row_net(self.batch.rows.get(),23),Decimal('200.00'))
        self.assertEqual(self.batch.uploads.count(),2)

    def test_repeated_own_import_does_not_duplicate_or_replace_price(self):
        self.row(supplier_lines=[], own={'name':'TEST CLIENT','amount':'999.00','email':'fill@example.test'})
        before = Booking.objects.count()
        import_own(self.batch,self.actor);import_own(self.batch,self.actor)
        self.assertEqual(Booking.objects.count(),before)
        self.booking.refresh_from_db();self.customer.refresh_from_db()
        self.assertEqual(self.booking.total_price_gross,Decimal('123'))
        self.assertEqual(self.customer.email,'fill@example.test')

    def test_shared_email_does_not_merge_different_people(self):
        for number,name in [('456','NEW FIRST'),('789','NEW SECOND')]:
            self.row(number=number,booking=None,own={'name':name,'amount':'100','email':'shared@example.test'},supplier_lines=[])
        import_own(self.batch,self.actor)
        self.assertEqual(Customer.objects.filter(email='shared@example.test').count(),2)

    def test_close_cancel_carry_and_no_payment_assumption(self):
        row=self.row(decision='CANCEL', note='Confirmed cancellation')
        close(self.batch,self.actor)
        self.booking.refresh_from_db();self.batch.refresh_from_db()
        self.assertEqual(self.booking.status,'CANCELLED')
        self.assertEqual(CustomerEvent.objects.filter(customer=self.customer,event_type='CANCELLATION').count(),1)
        self.assertEqual(self.batch.received,0)
        self.assertEqual(self.batch.employee_paid,0)
        self.assertEqual(totals(self.batch)['commission'],0)
        self.assertEqual(self.booking.total_price_gross,123) # preserved history

    def test_carry_excluded_and_created_next_month(self):
        self.row(decision='CARRY',note='Next month')
        close(self.batch,self.actor)
        following=Reconciliation.objects.get(month=date(2026,10,1),supplier=self.supplier)
        self.assertEqual(following.rows.get().number,'123')
        self.assertEqual(following.rows.get().supplier_lines,[])
        self.assertEqual(totals(self.batch)['commission'],0)

    def test_close_blocks_review_and_missing_net(self):
        row=self.row(decision='REVIEW')
        with self.assertRaises(ValidationError):close(self.batch,self.actor)
        row.decision='INCLUDE';row.supplier_lines=[{}];row.save()
        with self.assertRaises(ValidationError):close(self.batch,self.actor)

    def test_cancelled_booking_cannot_be_included(self):
        row=self.row();Booking.objects.filter(pk=self.booking.pk).update(status='CANCELLED')
        with self.assertRaises(ValidationError):decide(self.batch,row.pk,'INCLUDE','checked',self.actor)

    def test_employee_amount_never_used_as_commission_base(self):
        self.row()
        self.assertEqual(totals(self.batch)['commission'],Decimal('4.00'))

    def test_immutable_versions_and_permissions(self):
        self.row()
        first=checkpoint(self.batch,self.actor,'first')
        old=bytes(first.excel)
        self.batch.invoice_number='NEW';self.batch.save()
        checkpoint(self.batch,self.actor,'second')
        first.refresh_from_db();self.assertEqual(bytes(first.excel),old)
        url=reverse('admin:reconciliation_excel',args=[self.batch.pk])
        self.assertEqual(self.client.get(url).status_code,302)
        self.client.force_login(self.actor)
        response=self.client.get(url,{'version':first.pk})
        self.assertEqual(response.content,old)
        self.assertEqual(self.client.get(reverse('admin:reconciliation_workspace',args=[self.batch.pk])).status_code,200)

    def test_payments_require_invoice_and_dates(self):
        self.batch.received=10
        with self.assertRaises(ValidationError):self.batch.full_clean()
        self.batch.state='CLOSED';self.batch.invoice_number='TEST-INV';self.batch.invoice_amount=100;self.batch.received_date=date.today()
        self.batch.full_clean()

    def test_closed_upload_blocked(self):
        self.batch.state='CLOSED';self.batch.save()
        with self.assertRaises(ValidationError):stage(self.batch,'OWN','x',b'x',{},[],self.actor)

    def test_direct_orders_require_both_signals_and_exclude_employee_commission(self):
        self.batch.direct_orderer='TEST DIRECT';self.batch.direct_prefixes='W,H';self.batch.direct_rate=4;self.batch.save()
        lines=[{'number':'123','alternate_number':'W999','orderer':'TEST DIRECT','net':'100','row':2}]
        stage(self.batch,'SUPPLIER','direct.xlsx',b'direct',{},lines,self.actor)
        row=self.batch.rows.get();self.assertEqual(row.decision,'DIRECT')
        row.employee_data={'amount':'900'};row.save()
        result=totals(self.batch)
        self.assertEqual(result['commission'],0)
        self.assertEqual(result['direct_commission'],Decimal('4'))
        from .services import is_direct
        self.assertFalse(is_direct(self.batch,[dict(lines[0],orderer='OTHER')]))
        self.assertFalse(is_direct(self.batch,[dict(lines[0],alternate_number='999')]))

    def test_employee_replacement_removes_old_membership(self):
        self.row()
        stage(self.batch,'EMPLOYEE','new.xlsx',b'new-list',{},[{'number':'456','amount':'100','row':2}],self.actor)
        self.assertEqual(self.batch.rows.get(number='123').employee_data,{})
        self.assertEqual(totals(self.batch)['commission'],0)

    def test_prior_month_excludes_repeat_and_restores_own_excel(self):
        prior=self.row()
        october=Reconciliation.objects.create(month=date(2026,10,1),supplier=self.supplier,vat=23)
        book=Workbook();sheet=book.active
        sheet.append(['Reservation','Name','Amount']);sheet.append(['123','TEST CLIENT',123])
        content=BytesIO();book.save(content)
        mapping={'number':'A','name':'B','amount':'C'}
        extracted=read_rows(content.getvalue(),mapping)
        stage(october,'OWN','own.xlsx',content.getvalue(),mapping,extracted,self.actor)
        row=october.rows.get()
        self.assertEqual(row.decision,'PREVIOUS')
        self.assertEqual(row.previous_row,prior)
        annotated=load_workbook(BytesIO(bytes(october.revisions.first().own_excel)))
        self.assertEqual(annotated.active['A2'].fill.fgColor.rgb[-6:],'FFFF00')
        self.assertEqual(annotated.active['C2'].value,123)
        stage(october,'SUPPLIER','repeat.xlsx',b'repeat',{},[{'number':'123','net':'200','row':2}],self.actor)
        row.refresh_from_db();self.assertEqual(row.decision,'REVIEW')
        with self.assertRaises(ValidationError):decide(october,row.pk,'INCLUDE','verified',self.actor)
        row.employee_data={'amount':'999'};row.decision='INCLUDE';row.save()
        self.assertEqual(totals(october)['commission'],0)
        with self.assertRaises(ValidationError):close(october,self.actor)

    def test_history_does_not_cross_suppliers_or_mark_carry_as_counted(self):
        self.row(decision='CARRY')
        october=Reconciliation.objects.create(month=date(2026,10,1),supplier=self.supplier)
        stage(october,'OWN','carry.xlsx',b'carry',{},[{'number':'123','row':2}],self.actor)
        self.assertEqual(october.rows.get().decision,'REVIEW')
        other=Supplier.objects.create(supplier_code='OTHER',supplier_name='OTHER')
        other_batch=Reconciliation.objects.create(month=date(2026,10,1),supplier=other)
        stage(other_batch,'OWN','other.xlsx',b'other',{},[{'number':'123','row':2}],self.actor)
        self.assertIsNone(other_batch.rows.get().previous_row)

    def test_selective_import_preserves_time_and_is_idempotent(self):
        row=self.row(number='456',booking=None,supplier_lines=[],own={'name':'NEW CLIENT','amount':'100','pickup':'2026-10-11','pickup_raw':'11/10/26, 10:30, Airport'})
        excluded=self.row(number='789',booking=None,supplier_lines=[],own={'name':'OTHER CLIENT','amount':'100'})
        first=import_own(self.batch,self.actor,row_ids=[row.pk])
        second=import_own(self.batch,self.actor,row_ids=[row.pk])
        self.assertEqual(first['created'],1);self.assertEqual(second['created'],0)
        self.assertFalse(Booking.objects.filter(supplier_booking_number=excluded.number).exists())
        row.refresh_from_db()
        from django.utils import timezone
        self.assertEqual(timezone.localtime(row.booking.pickup_datetime).strftime('%H:%M'),'10:30')

    def test_gross_conversion_rounded_on_total(self):
        self.row(supplier_lines=[{'gross':'1.01'}])
        self.row(number='456',supplier_lines=[{'gross':'1.01'}])
        self.assertEqual(totals(self.batch)['net'],money(Decimal('2.02')/Decimal('1.23')))

    def test_correct_number_links_without_duplicate(self):
        row=self.row(number='MISSING-1',booking=None,issues=['Missing number'],decision='REVIEW')
        decide(self.batch,row.pk,'INCLUDE','Number checked',self.actor,corrected_number='RC123')
        row.refresh_from_db()
        self.assertEqual(row.number,'123')
        self.assertEqual(row.booking,self.booking)

    def test_workflow_post_and_revision_download(self):
        self.row()
        self.client.force_login(self.actor)
        url=reverse('admin:reconciliation_workspace',args=[self.batch.pk])
        response=self.client.post(url,{'action':'close'})
        self.assertEqual(response.status_code,302)
        self.batch.refresh_from_db();self.assertEqual(self.batch.state,'CLOSED')
        response=self.client.get(reverse('admin:reconciliation_excel',args=[self.batch.pk]))
        sheet=load_workbook(BytesIO(response.content)).get_sheet_by_name('Employee')
        self.assertEqual(sheet['A2'].value,'123')
        ordinary=get_user_model().objects.create_user('no-permission',password='test',is_staff=True)
        self.client.force_login(ordinary)
        self.assertEqual(self.client.get(url).status_code,403)
