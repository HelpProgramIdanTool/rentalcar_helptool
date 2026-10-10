import hashlib
import json
import re
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
from django.conf import settings
from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from bookings.models import Booking
from customers.models import Customer, CustomerEvent
from .models import Reconciliation, Upload, Row, Revision
from .parsing import identifier
from .history import prior_rows, annotate_own


def booking_index(supplier):
    result = {}
    for booking in Booking.objects.filter(supplier=supplier).select_related('customer'):
        if booking.supplier_booking_number:
            result.setdefault(identifier(booking.supplier_booking_number), []).append(booking)
    return result


def money(value):
    return Decimal(value).quantize(Decimal('.01'), rounding=ROUND_HALF_UP)


def row_net(row, vat, rounded=True):
    if not row.supplier_lines:
        return None
    total = Decimal(0)
    for line in row.supplier_lines:
        if line.get('net') is not None:
            total += Decimal(line['net'])
        elif line.get('gross') is not None and vat is not None:
            total += Decimal(line['gross']) / (1 + Decimal(vat) / 100)
        else:
            return None
    return money(total) if rounded else total


def totals(batch):
    earlier = {n for n,r in prior_rows(batch).items() if r.decision in ('INCLUDE','DIRECT')}
    net = Decimal(0)
    count = 0
    for row in batch.rows.filter(decision='INCLUDE'):
        if identifier(row.number) in earlier:
            continue
        if row.employee_data:
            value = row_net(row, batch.vat, rounded=False)
            if value is not None:
                net += value
                count += 1
    direct = [r for r in batch.rows.filter(decision='DIRECT') if identifier(r.number) not in earlier]
    direct_net = money(sum((row_net(r, batch.vat, rounded=False) or Decimal(0) for r in direct), Decimal(0)))
    return {'employee_orders': count, 'net': money(net), 'commission': money(money(net) * batch.rate / 100), 'direct_orders':len(direct), 'direct_net':direct_net, 'direct_commission':money(direct_net * batch.direct_rate / 100)}


def is_direct(batch, lines):
    prefixes = tuple(p.strip().upper() for p in batch.direct_prefixes.split(',') if p.strip())
    return bool(batch.direct_orderer.strip() and prefixes and lines and all(
        str(line.get('orderer','')).strip().casefold() == batch.direct_orderer.strip().casefold()
        and str(line.get('alternate_number') or line.get('number','')).upper().startswith(prefixes)
        for line in lines))


def require_open(batch):
    if batch.state != 'DRAFT':
        raise ValidationError('Сверка завершена. Сначала откройте её для исправления с указанием причины.')


@transaction.atomic
def stage(batch, kind, name, content, mapping, rows, actor):
    batch = Reconciliation.objects.select_for_update().get(pk=batch.pk)
    require_open(batch)
    digest = hashlib.sha256(content).hexdigest()
    upload, created = Upload.objects.get_or_create(reconciliation=batch, kind=kind, digest=digest,
        defaults={'name': Path(name).name, 'content': content, 'mapping': mapping})
    if not created:
        return False
    # Each upload of a statement is a replacement, not an addition. Multiple
    # invoice lines inside that statement remain separate evidence.
    if kind == 'SUPPLIER':
        for item in batch.rows.exclude(supplier_lines=[]):
            item.supplier_lines = []
            if item.decision in ('INCLUDE','DIRECT'):
                item.decision = 'REVIEW'
            item.save()
    if kind == 'EMPLOYEE':
        batch.rows.exclude(employee_data={}).update(employee_data={})
    known = booking_index(batch.supplier)
    history = prior_rows(batch)
    grouped = {}
    for data in rows:
        key = data['number'] or f'MISSING-{digest[:12]}-{data["row"]}'
        alternate = data.get('alternate_number','')
        if kind == 'SUPPLIER' and not known.get(key) and len(known.get(alternate, [])) == 1:
            key = identifier(known[alternate][0].supplier_booking_number)
        grouped.setdefault(key, []).append(data)
    for number, values in grouped.items():
        row, _ = Row.objects.get_or_create(reconciliation=batch, number=number)
        matches = known.get(identifier(number), [])
        row.booking = matches[0] if len(matches) == 1 else None
        issues = []
        if number.startswith('MISSING-'):
            issues.append('Missing reservation number: link to an existing booking or enter the correct number.')
        if len(matches) > 1:
            issues.append('Several database bookings have this supplier number; choose the correct booking.')
        if kind == 'SUPPLIER':
            row.supplier_lines = values
            if any(known.get(v.get('alternate_number')) and {b.pk for b in known[v['alternate_number']]} != {b.pk for b in matches} for v in values):
                issues.append('Primary and alternate numbers refer to different bookings; verify identity.')
            if len(values) > 1:
                issues.append('Multiple statement lines: confirm separate invoices before including.')
        else:
            field = 'own' if kind == 'OWN' else 'employee_data'
            previous = getattr(row, field)
            if len(values) > 1:
                issues.append('Repeated source number: check all source rows before confirming.')
            setattr(row, field, values[0])
            if previous and previous != values[0]:
                issues.append('Source changed since previous upload. Previous values are retained in Excel versions.')
        if row.booking and row.own:
            old = row.booking.customer_name_snapshot or str(row.booking.customer)
            if row.own.get('name') and row.own['name'].casefold() != old.casefold():
                issues.append('Name differs from database; existing customer will not be replaced.')
            if row.own.get('amount') is not None and Decimal(row.own['amount']) != row.booking.total_price_gross:
                issues.append('Own amount differs from database; existing booking price will not be overwritten.')
        row.issues = list(dict.fromkeys(row.issues + issues))
        row.decision = 'REVIEW'
        if is_direct(batch, row.supplier_lines) and row_net(row, batch.vat) is not None:
            row.decision = 'DIRECT'
        elif not row.issues and row.booking and row.supplier_lines and row.booking.status != 'CANCELLED' and not row.own.get('cancelled') and row_net(row, batch.vat) is not None:
            row.decision = 'INCLUDE'
        elif not row.supplier_lines and not row.employee_data and row.own.get('return') and row.own['return'][:7] != batch.month.strftime('%Y-%m'):
            row.decision = 'EXCLUDE'
            row.note = 'Own historical record outside this month; available for filling missing database records.'
        previous = history.get(identifier(number))
        if previous:
            row.previous_row = previous
            if previous.decision in ('INCLUDE','DIRECT'):
                row.decision = 'REVIEW' if row.supplier_lines else 'PREVIOUS'
                row.note = f'Previously reconciled: {previous.reconciliation.month:%Y-%m}. Excluded from new commission.'
                if row.supplier_lines:
                    row.issues = list(dict.fromkeys(row.issues + ['Repeated supplier reservation from a prior month. Check duplicate versus additional charge; no automatic commission.']))
            elif previous.decision=='CANCEL' and row.booking and row.booking.status=='CANCELLED' and not row.supplier_lines:
                row.decision='CANCEL'
                row.note=f'Cancellation preserved from {previous.reconciliation.month:%Y-%m}.'
            elif previous.decision=='CARRY' and not row.supplier_lines:
                row.decision='REVIEW'
                row.note=f'Carry-forward from {previous.reconciliation.month:%Y-%m}; awaiting current supplier statement.'
        row.save()
    checkpoint(batch, actor, f'Uploaded {kind}: {Path(name).name}')
    return True


@transaction.atomic
def import_own(batch, actor, row_ids=None):
    batch = Reconciliation.objects.select_for_update().get(pk=batch.pk)
    require_open(batch)
    counts = {'created': 0, 'filled': 0, 'unchanged': 0, 'review': 0}
    type(batch.supplier).objects.select_for_update().get(pk=batch.supplier_id)
    known = booking_index(batch.supplier)
    selected = batch.rows.exclude(own={})
    if row_ids is not None:
        selected = selected.filter(pk__in=row_ids)
    for row in selected:
        data = row.own
        if row.number.startswith('MISSING-') or (row.issues and (not row.note or row.decision == 'EXCLUDE')):
            counts['review'] += 1
            continue
        matches = known.get(identifier(row.number), [])
        if len(matches) > 1:
            counts['review'] += 1
            continue
        booking = matches[0] if matches else None
        if not booking:
            name = data.get('name', '').replace('\u200b','').strip()
            if not name:
                counts['review'] += 1
                continue
            # Email alone is never a customer identity.
            customers = Customer.objects.filter(full_name_latin__iexact=name, email__iexact=data.get('email', '')) if data.get('email') else Customer.objects.none()
            customer = customers.first() if customers.count() == 1 else None
            if not customer:
                first, _, last = name.partition(' ')
                customer = Customer.objects.create(first_name=first[:100], last_name=last[:100], full_name_latin=name[:200], email=data.get('email','')[:254], phone_1=data.get('phone','')[:30])
            booking = Booking(customer=customer, supplier=batch.supplier, supplier_booking_number=row.number,
                status='DRAFT', source_quote_snapshot={'historical_import': {'reconciliation_id': batch.pk, 'source': data}},
                total_price_gross=Decimal(data.get('amount') or 0), vehicle_price_gross=Decimal(data.get('amount') or 0))
            booking.save()
            known.setdefault(identifier(row.number), []).append(booking)
            counts['created'] += 1
        else:
            counts['unchanged'] += 1
        updates = {}
        for field, key in [('pickup_location_text','pickup_raw'), ('return_location_text','return_raw')]:
            if not getattr(booking, field) and data.get(key):
                updates[field] = data[key][:300]
        for field, key in [('pickup_datetime','pickup'), ('return_datetime','return')]:
            if not getattr(booking, field) and data.get(key):
                value = datetime.fromisoformat(data[key])
                time_match = re.search(r'\b([01]?\d|2[0-3]):([0-5]\d)\b',data.get(key+'_raw',''))
                if time_match:
                    value = value.replace(hour=int(time_match[1]),minute=int(time_match[2]))
                updates[field] = timezone.make_aware(value)
        if updates:
            Booking.objects.filter(pk=booking.pk).update(**updates)
            counts['filled'] += 1
        # Never overwrite populated customer contacts or names.
        customer = booking.customer
        for field, key in [('email','email'), ('phone_1','phone')]:
            if not getattr(customer, field) and data.get(key):
                setattr(customer, field, data[key][:254 if field == 'email' else 30])
                customer.save(update_fields=[field])
        row.booking = booking
        row.save(update_fields=['booking'])
    checkpoint(batch, actor, 'Imported missing own records: ' + str(counts))
    return counts


@transaction.atomic
def decide(batch, row_id, decision, note, actor, booking_id=None, corrected_number=''):
    batch = Reconciliation.objects.select_for_update().get(pk=batch.pk)
    require_open(batch)
    row = batch.rows.select_for_update().get(pk=row_id)
    if corrected_number.strip():
        number = identifier(corrected_number)
        if batch.rows.exclude(pk=row.pk).filter(number=number).exists():
            raise ValidationError('Этот номер уже есть в сверке. Не создавайте дубликат.')
        if row.booking and identifier(row.booking.supplier_booking_number) != number:
            raise ValidationError('Сначала выберите правильный заказ в базе.')
        row.number = number
        matches = booking_index(batch.supplier).get(number, [])
        if len(matches) == 1:
            row.booking = matches[0]
    if decision not in dict(Row._meta.get_field('decision').choices):
        raise ValidationError('Неизвестное решение.')
    previous = prior_rows(batch).get(identifier(row.number))
    if decision in ('INCLUDE','DIRECT') and previous and previous.decision in ('INCLUDE','DIRECT'):
        raise ValidationError('Этот заказ уже учтён ранее. Повторную сумму нельзя включить как новый заказ; нужна отдельная проверка корректировки.')
    if decision=='PREVIOUS' and (not previous or previous.decision not in ('INCLUDE','DIRECT')):
        raise ValidationError('Нет подтверждённой предыдущей сверки для этого номера.')
    if booking_id:
        booking = Booking.objects.get(pk=booking_id, supplier=batch.supplier)
        if batch.rows.exclude(pk=row.pk).filter(number=booking.supplier_booking_number).exists():
            raise ValidationError('Этот заказ уже есть в сверке. Устраните дубль.')
        row.booking = booking
        row.number = booking.supplier_booking_number
    if not note.strip() and (row.issues or decision in ('CANCEL','CARRY','EXCLUDE')):
        raise ValidationError('Укажите причину решения.')
    if decision == 'INCLUDE':
        if not row.booking or not row.supplier_lines:
            raise ValidationError('Нужны заказ в базе и строка поставщика.')
        if row.booking.status == 'CANCELLED' or row.own.get('cancelled'):
            raise ValidationError('Отменённый заказ нельзя включить. Сначала разберите отмену в карточке заказа / исходном файле.')
        if row_net(row, batch.vat) is None:
            raise ValidationError('Нет суммы нетто: укажите подтверждённый НДС или загрузите нетто поставщика.')
    if decision == 'DIRECT' and (not is_direct(batch, row.supplier_lines) or row_net(row,batch.vat) is None):
        raise ValidationError('Прямой заказ должен соответствовать заказывающему и префиксу номера; нужна сумма нетто.')
    if decision == 'CANCEL' and not row.booking:
        raise ValidationError('Сначала свяжите отмену с заказом в базе.')
    row.decision, row.note = decision, note
    row.save()
    checkpoint(batch, actor, f'{row.number}: {decision} — {note}'[:255])


@transaction.atomic
def close(batch, actor):
    batch = Reconciliation.objects.select_for_update().get(pk=batch.pk)
    require_open(batch)
    if not batch.rows.exists() or batch.rows.filter(decision='REVIEW').exists():
        raise ValidationError('Сначала разберите все строки со статусом «Проверить».')
    history = prior_rows(batch)
    for row in batch.rows.select_related('booking'):
        previous = history.get(identifier(row.number))
        if row.decision in ('INCLUDE','DIRECT') and previous and previous.decision in ('INCLUDE','DIRECT'):
            raise ValidationError(f'{row.number} уже учтён ранее. Разберите повторное начисление.')
        if row.decision=='PREVIOUS' and (not previous or previous.decision not in ('INCLUDE','DIRECT')):
            raise ValidationError(f'История {row.number} изменилась; проверьте прежнюю сверку.')
        if row.decision == 'DIRECT' and (not is_direct(batch,row.supplier_lines) or row_net(row,batch.vat) is None):
            raise ValidationError(f'Проверьте правило прямого заказа {row.number}.')
        if row.decision == 'INCLUDE' and (not row.booking or not row.supplier_lines or row.booking.status == 'CANCELLED' or row_net(row, batch.vat) is None):
            raise ValidationError(f'Не подтверждён заказ {row.number}.')
        if row.decision == 'CANCEL':
            if not row.booking:
                raise ValidationError(f'Отмена {row.number} не связана с клиентом.')
            booking = row.booking
            snapshot = dict(booking.source_quote_snapshot)
            snapshot.setdefault('reconciliation_original_price', str(booking.total_price_gross))
            Booking.objects.filter(pk=booking.pk).update(status='CANCELLED', source_quote_snapshot=snapshot)
            if not CustomerEvent.objects.filter(customer=booking.customer, supplier=batch.supplier, event_type='CANCELLATION', title__contains=row.number).exists():
                CustomerEvent.objects.create(customer=booking.customer, supplier=batch.supplier,
                    event_type='CANCELLATION', title=f'Отмена заказа {row.number}',
                    event_date=date.today(), description=f'Сверка {batch.month:%Y-%m}. {row.note}', severity='INFO')
        if row.decision == 'CARRY':
            next_month = date(batch.month.year + (batch.month.month == 12), batch.month.month % 12 + 1, 1)
            following, _ = Reconciliation.objects.get_or_create(month=next_month, supplier=batch.supplier,
                defaults={'rate': batch.rate, 'vat': batch.vat, 'employee': batch.employee, 'direct_orderer':batch.direct_orderer, 'direct_prefixes':batch.direct_prefixes, 'direct_rate':batch.direct_rate})
            require_open(following)
            carried, created = Row.objects.get_or_create(reconciliation=following, number=row.number,
                defaults={'booking':row.booking, 'own':row.own, 'employee_data':row.employee_data, 'note':f'Перенос из {batch.month:%Y-%m}'})
            if not created and carried.booking_id != row.booking_id:
                raise ValidationError(f'Конфликт переноса {row.number}.')
    batch.state = 'CLOSED'
    batch.save()
    return checkpoint(batch, actor, 'Сверка завершена; внесите номер фактуры')


def workbook(snapshot):
    book = Workbook()
    summary = book.active
    summary.title = 'Summary'
    for key, value in snapshot['summary'].items():
        if key in ('Employee rate (%)','VAT (%)','Employee net turnover','Employee commission','Invoice amount','Received','Employee paid','Total employee net','Total employee commission','Direct supplier net','My direct commission','Direct commission rate (%)') and value != '':
            value = float(Decimal(value))
        summary.append([key, value])
    headings = ['Reservation', 'Customer', 'Own amount', 'Supplier gross', 'Supplier net', 'Employee order', 'Decision', 'Description']
    for title, selected in [('Orders', snapshot['rows']), ('Cancellations', [r for r in snapshot['rows'] if r['decision']=='CANCEL']), ('Carry forward', [r for r in snapshot['rows'] if r['decision']=='CARRY']), ('Employee', [r for r in snapshot['rows'] if r['employee']]), ('Review', [r for r in snapshot['rows'] if r['decision']=='REVIEW'])]:
        sheet = book.create_sheet(title)
        sheet.append(headings)
        for row in selected:
            values = [row[k] for k in ('number','name','own_amount','gross','net','employee','decision','note')]
            for index in (2,3,4):
                if values[index] not in ('', None):
                    values[index] = float(Decimal(values[index]))
            sheet.append(values)
            fill = 'FDE2E2' if row['decision']=='CANCEL' else 'FFF3B0' if row['decision']=='INCLUDE' else 'FFFFFF'
            for cell in sheet[sheet.max_row]:
                cell.fill = PatternFill('solid', fgColor=fill)
    for sheet in book:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='17365D')
        for row in sheet:
            for cell in row:
                # Never execute source text as an Excel formula.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                elif isinstance(cell.value, (float, Decimal)):
                    cell.number_format = '#,##0.00'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 28 if col[0].column < 8 else 65
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def checkpoint(batch, actor, reason):
    calculation = totals(batch)
    summary = {'Period': str(batch.month), 'Supplier': str(batch.supplier), 'State': batch.state,
        'Direct supplier orders':calculation['direct_orders'], 'Direct supplier net':str(calculation['direct_net']), 'My direct commission':str(calculation['direct_commission']), 'Direct commission rate (%)':str(batch.direct_rate),
        'Employee rate (%)': str(batch.rate), 'VAT (%)': str(batch.vat) if batch.vat is not None else '',
        'Employee net turnover': str(calculation['net']), 'Employee commission': str(calculation['commission']),
        'Invoice number': batch.invoice_number, 'Invoice date': str(batch.invoice_date or ''),
        'Invoice amount': str(batch.invoice_amount or ''), 'Received': str(batch.received),
        'Received date': str(batch.received_date or ''), 'Employee paid': str(batch.employee_paid),
        'Employee payment date': str(batch.employee_paid_date or '')}
    rows = []
    for row in batch.rows.select_related('booking'):
        gross_values = [x.get('gross') for x in row.supplier_lines]
        rows.append({'number':row.number, 'name':row.own.get('name') or (row.booking.customer_name_snapshot if row.booking else ''),
            'own_amount':row.own.get('amount') or '', 'gross':str(sum(Decimal(x) for x in gross_values)) if gross_values and all(x is not None for x in gross_values) else '',
            'net':str(row_net(row, batch.vat)) if row_net(row, batch.vat) is not None else '', 'employee':bool(row.employee_data),
            'decision':row.decision, 'note':row.note + ' ' + '; '.join(row.issues),
            'evidence':{'own':row.own, 'supplier':row.supplier_lines, 'employee':row.employee_data}})
    snapshot = {'summary':summary, 'rows':rows}
    content = workbook(snapshot)
    own_upload = batch.uploads.filter(kind='OWN').order_by('-pk').first()
    own_content = annotate_own(own_upload,batch) if own_upload else None
    revision = Revision.objects.create(reconciliation=batch, actor=str(actor), reason=reason[:255], snapshot=snapshot,
        excel=content, own_excel=own_content, digest=hashlib.sha256(content).hexdigest())
    # Separate private Excel files survive an accidental deletion of a DB row.
    # An off-device backup directory may be configured by the operator.
    root = Path(getattr(settings, 'RECONCILIATION_EXPORT_ROOT', settings.BASE_DIR / 'private_uploads/monthly_exports'))
    directory = root / batch.month.strftime('%Y-%m') / str(batch.supplier_id)
    directory.mkdir(parents=True, exist_ok=True)
    (directory / f'reconciliation_{batch.pk}_v{revision.pk}.xlsx').write_bytes(content)
    if own_content:
        (directory / f'own_orders_{batch.pk}_v{revision.pk}.xlsx').write_bytes(own_content)
    return revision
