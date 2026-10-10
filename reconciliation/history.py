"""Previous accepted decisions; invoice payments are deliberately independent."""
from collections import Counter
from copy import copy
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.cell import column_index_from_string
from .models import Row
from .parsing import identifier


def prior_rows(batch):
    result = {}
    for row in Row.objects.filter(reconciliation__supplier=batch.supplier,
            reconciliation__month__lt=batch.month,
            decision__in=['INCLUDE','DIRECT','CANCEL','CARRY']).select_related('reconciliation').order_by('-reconciliation__month','-pk'):
        if row.decision in ('INCLUDE','DIRECT') and not row.supplier_lines:
            continue
        result.setdefault(identifier(row.number),row)
    return result


def annotate_own(upload, batch):
    mapping = upload.mapping
    if not mapping.get('number'):
        return None
    book = load_workbook(BytesIO(bytes(upload.content)))
    sheet = book[mapping['_sheet']] if mapping.get('_sheet') else book.active
    number_col = column_index_from_string(mapping['number'])
    amount_col = column_index_from_string(mapping['amount']) if mapping.get('amount') else None
    first = int(mapping.get('_first_row',2))
    numbers = {index:identifier(sheet.cell(index,number_col).value) for index in range(first,sheet.max_row+1)}
    counts = Counter(numbers.values())
    history = prior_rows(batch)
    status_col = sheet.max_column + 1
    sheet.cell(1,status_col,'Reconciliation history (not invoice payment)')
    sheet.column_dimensions[sheet.cell(1,status_col).column_letter].width=65
    for index,number in numbers.items():
        if not number or counts[number] != 1 or number not in history:
            continue
        previous = history[number]
        period = previous.reconciliation.month.strftime('%Y-%m')
        if previous.decision in ('INCLUDE','DIRECT'):
            message=f'Reconciled: {period}. Do not count again. Invoice payment tracked separately.'
            fill='FFFF00'
        elif previous.decision=='CANCEL':
            message=f'Cancelled: confirmed history ({period}). Excluded from commission.'
            fill='FFC7CE'
            if amount_col:
                sheet.cell(index,amount_col).value=None
        else:
            message=f'Carry forward from {period}; awaiting supplier settlement. Not previously counted.'
            fill=None
        if fill:
            for col in range(1,max(number_col,amount_col or 1)+1):
                cell=sheet.cell(index,col)
                cell.fill=PatternFill('solid',fgColor=fill)
                if previous.decision=='CANCEL':
                    font=copy(cell.font);font.color='9C0006';cell.font=font
        sheet.cell(index,status_col,message)
    stream=BytesIO();book.save(stream);book.close()
    return stream.getvalue()
